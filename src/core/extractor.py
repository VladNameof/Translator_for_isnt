"""
Deep AST Text Extractor Module for PowerPoint Presentations
Recursively extracts text from Slide Masters, Groups, Tables, and Shapes.
"""

from typing import List, Dict, Any
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from src.core.normalizer import TextNormalizer

class ExtractedItem:
    def __init__(self, slide_idx: Any, path: str, text: str, shape_type: str):
        self.slide = slide_idx
        self.path = path
        self.text = text
        self.shape_type = shape_type
        self.normalized_text = TextNormalizer.normalize(text)
        self.is_translatable = TextNormalizer.is_translatable(text)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "slide": self.slide,
            "path": self.path,
            "text": self.text,
            "normalized": self.normalized_text,
            "shape_type": self.shape_type,
            "is_translatable": self.is_translatable
        }


class PPTXExtractor:
    @staticmethod
    def extract_from_presentation(pptx_path: str) -> List[ExtractedItem]:
        """
        Recursively extracts all text elements from a PPTX file.
        """
        prs = Presentation(pptx_path)
        items: List[ExtractedItem] = []

        def process_shape(shape, path_prefix: str, slide_id: Any):
            # 1. Grouped Shapes (Recursive traversal)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for idx, child in enumerate(shape.shapes, start=1):
                    child_name = getattr(child, 'name', f"Child_{idx}")
                    process_shape(child, f"{path_prefix}/Group[{shape.name}]/{child_name}", slide_id)

            # 2. Tables
            elif shape.has_table:
                for r_idx, row in enumerate(shape.table.rows, start=1):
                    for c_idx, cell in enumerate(row.cells, start=1):
                        if cell.text_frame and cell.text_frame.text:
                            raw = cell.text_frame.text
                            if raw.strip():
                                items.append(ExtractedItem(
                                    slide_idx=slide_id,
                                    path=f"{path_prefix}/Table[{shape.name}]/Cell({r_idx},{c_idx})",
                                    text=raw,
                                    shape_type="TableCell"
                                ))

            # 3. Standard Shapes with TextFrame
            elif shape.has_text_frame:
                raw = shape.text_frame.text
                if raw and raw.strip():
                    items.append(ExtractedItem(
                        slide_idx=slide_id,
                        path=f"{path_prefix}/{shape.name}",
                        text=raw,
                        shape_type="TextFrame"
                    ))

        # 1. Slide Masters
        for m_idx, master in enumerate(prs.slide_masters, start=1):
            for s_idx, shape in enumerate(master.shapes, start=1):
                process_shape(shape, f"Master_{m_idx}/Shape_{s_idx}", "Master")

        # 2. Slides
        for s_idx, slide in enumerate(prs.slides, start=1):
            for sh_idx, shape in enumerate(slide.shapes, start=1):
                process_shape(shape, f"Slide_{s_idx}/Shape_{sh_idx}", s_idx)

        return items

    @staticmethod
    def get_unique_translatable_strings(items: List[ExtractedItem]) -> List[str]:
        """
        Filters and returns unique translatable strings from extracted items.
        """
        seen = set()
        unique_list = []
        for item in items:
            if item.is_translatable:
                norm = item.normalized_text
                if norm not in seen:
                    seen.add(norm)
                    unique_list.append(item.text)
        return unique_list
